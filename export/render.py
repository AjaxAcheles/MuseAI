"""Render a finished story (Markdown source) to a downloadable file.

Stories are stored as light Markdown: a single '# Title' heading, prose
paragraphs, and '* * *' scene breaks. We render that to md / txt / docx / pdf.
"""
from __future__ import annotations

import io
import re

MIME = {
    "md": "text/markdown",
    "txt": "text/plain",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


def render(content_md: str, fmt: str) -> tuple[bytes, str]:
    """Return (bytes, mimetype) for the requested format."""
    fmt = fmt.lower()
    if fmt == "md":
        return content_md.encode("utf-8"), MIME["md"]
    title, blocks = _parse(content_md)
    if fmt == "txt":
        return _to_txt(title, blocks).encode("utf-8"), MIME["txt"]
    if fmt == "docx":
        return _to_docx(title, blocks), MIME["docx"]
    if fmt == "pdf":
        return _to_pdf(title, blocks), MIME["pdf"]
    raise ValueError(f"Unsupported export format: {fmt}")


def _parse(content_md: str) -> tuple[str, list[tuple[str, str]]]:
    """Split into a title and a list of (kind, text) blocks where kind is
    'break' or 'para'."""
    lines = content_md.splitlines()
    title = "Untitled"
    body_start = 0
    for i, ln in enumerate(lines):
        if ln.startswith("# "):
            title = ln[2:].strip()
            body_start = i + 1
            break
    body = "\n".join(lines[body_start:]).strip()
    blocks: list[tuple[str, str]] = []
    for chunk in re.split(r"\n\s*\n", body):
        chunk = chunk.strip()
        if not chunk:
            continue
        if chunk in ("* * *", "***", "---"):
            blocks.append(("break", ""))
        else:
            blocks.append(("para", " ".join(chunk.split("\n"))))
    return title, blocks


def _to_txt(title: str, blocks: list[tuple[str, str]]) -> str:
    out = [title, "=" * len(title), ""]
    for kind, text in blocks:
        out.append("* * *" if kind == "break" else text)
        out.append("")
    return "\n".join(out)


def _to_docx(title: str, blocks: list[tuple[str, str]]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    doc.add_heading(title, level=0)
    for kind, text in blocks:
        if kind == "break":
            p = doc.add_paragraph("* * *")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _to_pdf(title: str, blocks: list[tuple[str, str]]) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import inch
    from xml.sax.saxutils import escape

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=letter,
        leftMargin=1 * inch, rightMargin=1 * inch, topMargin=1 * inch, bottomMargin=1 * inch,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("MuseTitle", parent=styles["Title"], fontName="Times-Bold")
    body_style = ParagraphStyle(
        "MuseBody", parent=styles["BodyText"], fontName="Times-Roman",
        fontSize=12, leading=18, alignment=TA_JUSTIFY, firstLineIndent=18, spaceAfter=2,
    )
    break_style = ParagraphStyle("MuseBreak", parent=styles["BodyText"], alignment=TA_CENTER, spaceBefore=10, spaceAfter=10)

    flow = [Paragraph(escape(title), title_style), Spacer(1, 18)]
    for kind, text in blocks:
        if kind == "break":
            flow.append(Paragraph("* * *", break_style))
        else:
            flow.append(Paragraph(escape(text), body_style))
    doc.build(flow)
    return buf.getvalue()
